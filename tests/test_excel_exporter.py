from openpyxl import load_workbook

from exporters.excel_exporter import ExcelExporter


def test_excel_exporter_creates_file(tmp_path):
    output_path = tmp_path / "businesses.xlsx"

    data = [
        {
            "name": "Pizza Sara",
            "city": "Mashhad",
            "phone": "09121234567",
        },
    ]

    exporter = ExcelExporter()

    result = exporter.export(data, output_path)

    assert result == output_path
    assert output_path.exists()


def test_excel_exporter_writes_headers_and_rows(tmp_path):
    output_path = tmp_path / "businesses.xlsx"

    data = [
        {
            "name": "Pizza Sara",
            "city": "Mashhad",
            "phone": "09121234567",
        },
        {
            "name": "Fast Food Center",
            "city": "Mashhad",
            "phone": "09129876543",
        },
    ]

    ExcelExporter().export(data, output_path)

    workbook = load_workbook(output_path)
    worksheet = workbook["Businesses"]

    headers = [
        cell.value
        for cell in worksheet[1]
    ]

    assert headers == [
        "Name",
        "City",
        "Phone",
    ]

    assert worksheet.cell(row=2, column=1).value == "Pizza Sara"
    assert worksheet.cell(row=2, column=2).value == "Mashhad"
    assert worksheet.cell(row=2, column=3).value == "09121234567"

    assert worksheet.cell(row=3, column=1).value == "Fast Food Center"
    assert worksheet.cell(row=3, column=2).value == "Mashhad"
    assert worksheet.cell(row=3, column=3).value == "09129876543"

    workbook.close()


def test_excel_exporter_uses_standard_business_column_order(
    tmp_path,
):
    output_path = tmp_path / "businesses.xlsx"

    data = [
        {
            "phone": "09121234567",
            "name": "Pizza Sara",
            "rating": 4.5,
            "city": "Mashhad",
        },
    ]

    ExcelExporter().export(data, output_path)

    workbook = load_workbook(output_path)
    worksheet = workbook["Businesses"]

    headers = [
        cell.value
        for cell in worksheet[1]
    ]

    assert headers == [
        "Name",
        "City",
        "Phone",
        "Rating",
    ]

    workbook.close()


def test_excel_exporter_preserves_persian_text(tmp_path):
    output_path = tmp_path / "businesses.xlsx"

    data = [
        {
            "name": "پیتزا سارا",
            "city": "مشهد",
            "address": "بلوار سجاد",
        },
    ]

    ExcelExporter().export(data, output_path)

    workbook = load_workbook(output_path)
    worksheet = workbook["Businesses"]

    assert worksheet.cell(row=2, column=1).value == "پیتزا سارا"
    assert worksheet.cell(row=2, column=2).value == "بلوار سجاد"
    assert worksheet.cell(row=2, column=3).value == "مشهد"
    
    workbook.close()


def test_excel_exporter_supports_string_output_path(tmp_path):
    output_path = tmp_path / "businesses.xlsx"

    data = [
        {
            "name": "Pizza Sara",
        },
    ]

    result = ExcelExporter().export(
        data,
        str(output_path),
    )

    assert result == output_path
    assert output_path.exists()


def test_excel_exporter_creates_parent_directory(tmp_path):
    output_path = (
        tmp_path
        / "exports"
        / "businesses.xlsx"
    )

    data = [
        {
            "name": "Pizza Sara",
        },
    ]

    ExcelExporter().export(data, output_path)

    assert output_path.exists()


def test_excel_exporter_handles_empty_data(tmp_path):
    output_path = tmp_path / "empty.xlsx"

    result = ExcelExporter().export([], output_path)

    assert result == output_path
    assert output_path.exists()

    workbook = load_workbook(output_path)

    assert "Businesses" in workbook.sheetnames
    assert workbook["Businesses"].max_row == 1
    assert workbook["Businesses"].max_column == 1
    assert workbook["Businesses"].cell(row=1, column=1).value is None

    workbook.close()


def test_excel_exporter_freezes_header_and_enables_filter(tmp_path):
    output_path = tmp_path / "businesses.xlsx"

    data = [
        {
            "name": "Pizza Sara",
            "city": "Mashhad",
            "phone": "09121234567",
        },
    ]

    ExcelExporter().export(data, output_path)

    workbook = load_workbook(output_path)
    worksheet = workbook["Businesses"]

    assert worksheet.freeze_panes == "A2"
    assert worksheet.auto_filter.ref == "A1:C2"

    workbook.close()