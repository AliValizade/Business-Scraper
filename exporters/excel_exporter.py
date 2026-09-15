from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from exporters.base import BaseExporter


class ExcelExporter(BaseExporter):
    """Export business data to an Excel workbook."""

    DEFAULT_COLUMNS = [
        "id",
        "name",
        "category",
        "address",
        "city",
        "phone",
        "website",
        "instagram",
        "rating",
        "reviews_count",
        "latitude",
        "longitude",
        "google_maps_url",
        "source",
        "source_id",
        "search_keyword",
        "source_url",
        "scraped_at",
        "created_at",
        "updated_at",
    ]

    COLUMN_HEADERS = {
        "id": "ID",
        "name": "Name",
        "category": "Category",
        "address": "Address",
        "city": "City",
        "phone": "Phone",
        "website": "Website",
        "instagram": "Instagram",
        "rating": "Rating",
        "reviews_count": "Reviews Count",
        "latitude": "Latitude",
        "longitude": "Longitude",
        "google_maps_url": "Google Maps URL",
        "source": "Source",
        "source_id": "Source ID",
        "search_keyword": "Search Keyword",
        "source_url": "Source URL",
        "scraped_at": "Scraped At",
        "created_at": "Created At",
        "updated_at": "Updated At",
    }

    def export(
        self,
        data: Iterable[dict[str, Any]],
        output_path: str | Path,
    ) -> Path:
        output_path = Path(output_path)

        rows = list(data)

        output_path.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Businesses"

        if not rows:
            workbook.save(output_path)
            return output_path

        columns = self._get_columns(rows)

        headers = [
            self.COLUMN_HEADERS.get(column, column)
            for column in columns
        ]

        worksheet.append(headers)

        for row in rows:
            worksheet.append(
                [
                    row.get(column)
                    for column in columns
                ]
            )

        self._adjust_column_widths(
            worksheet,
            len(columns),
        )

        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        workbook.save(output_path)

        return output_path

    @classmethod
    def _get_columns(
        cls,
        rows: list[dict[str, Any]],
    ) -> list[str]:
        columns = []

        for column in cls.DEFAULT_COLUMNS:
            if any(column in row for row in rows):
                columns.append(column)

        for row in rows:
            for key in row:
                if key not in columns:
                    columns.append(key)

        return columns

    @staticmethod
    def _adjust_column_widths(
        worksheet,
        column_count: int,
    ) -> None:
        for column_index in range(1, column_count + 1):
            column_letter = get_column_letter(column_index)

            max_length = 0

            for cell in worksheet[column_letter]:
                if cell.value is not None:
                    max_length = max(
                        max_length,
                        len(str(cell.value)),
                    )

            worksheet.column_dimensions[
                column_letter
            ].width = min(max_length + 2, 60)